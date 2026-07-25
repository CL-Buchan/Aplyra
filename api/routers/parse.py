import io

from fastapi import APIRouter, HTTPException, UploadFile
from pypdf import PdfReader
from docx import Document

from models import ParsedDocument

router = APIRouter()

PDF_TYPE = "application/pdf"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@router.post("/parse/document", response_model=ParsedDocument)
async def parse_document(file: UploadFile):
    if file.content_type not in (PDF_TYPE, DOCX_TYPE):
        raise HTTPException(
            status_code=415,
            detail=f"Unsupported file type: {file.content_type}. Upload a PDF or .docx.",
        )

    contents = await file.read()

    if file.content_type == PDF_TYPE:
        reader = PdfReader(io.BytesIO(contents))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    else:
        reader = Document(io.BytesIO(contents))
        text = "\n".join(para.text for para in reader.paragraphs)

    if not text.strip():
        raise HTTPException(status_code=422, detail="Could not extract any text from the document.")

    return ParsedDocument(filename=file.filename, text=text)
